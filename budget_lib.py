#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
预算读取模块：从《预算分解.xlsx》读取部门级日预算，提供 单日 / MTD累计 的未税预算。

口径（已对 2026-09 逐部门勾稽验证，差异=0.0%）：
  - 日预算表（9月/10月 sheet）中 部门_线上 / 部门_POS 为【含税成交】口径。
  - 日参数 sheet 的 税比 = 当日 含税/未税 系数（门店级，逐日不同）。
  - 部门当日未税预算 = (线上 + POS) / 税比；POS未税预算 = POS / 税比。
  - 单日Σ、MTD累计Σ 均与《部门月预算(未税)》完全吻合。
  - 末行"合计"行自动跳过，避免重复累加。

对外接口：
  BudgetReader(path)
    .departments()                 -> 预算文件里的部门名单
    .daily_pretax(date)            -> {dept: {'online':..,'pos':..,'total':..}} 当日未税
    .mtd_pretax(date)              -> {dept: {...}} 本月1日至 date 累计未税
    .has_date(date)                -> 该日是否有预算
  date 统一用 'YYYYMMDD' 字符串。
"""
import io, sys
from openpyxl import load_workbook

DEPTS = ['标准品部','肉类禽蛋部','日配部','水果部','蔬菜部','海鲜水产部','烘焙部',
         '中餐部','餐饮联营商部','海鲜岛','社区服务','行政部','其他部门']

def _norm(d):
    """把 '2026-09-01' / '20260901' / datetime 归一成 'YYYYMMDD'。"""
    s = str(d)[:10].replace('-', '')
    if len(s) >= 8 and s[:8].isdigit():
        return s[:8]
    return None

class BudgetReader:
    def __init__(self, path):
        wb = load_workbook(path, data_only=True)
        self.daily = {}      # date -> dept -> {'online_tax','pos_tax'}
        self.tax_ratio = {}  # date -> ratio
        self._parse_daily_sheets(wb)
        self._parse_tax(wb)

    def _header_map(self, ws):
        return {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column+1)}

    def _parse_daily_sheets(self, wb):
        for ws in wb.worksheets:
            h = self._header_map(ws)
            if '日期' not in h or '线上小计' not in h:
                continue  # 非日预算表
            for r in range(2, ws.max_row+1):
                raw_date = ws.cell(row=r, column=h['日期']).value
                if raw_date is None or str(raw_date).strip() == '合计':
                    continue
                d = _norm(raw_date)
                if not d:
                    continue
                rec = {}
                for dept in DEPTS:
                    on = ws.cell(row=r, column=h[f'{dept}_线上']).value if f'{dept}_线上' in h else None
                    po = ws.cell(row=r, column=h[f'{dept}_POS']).value if f'{dept}_POS' in h else None
                    rec[dept] = {'online_tax': float(on or 0), 'pos_tax': float(po or 0)}
                self.daily[d] = rec

    def _parse_tax(self, wb):
        if '日参数' not in wb.sheetnames:
            return
        ws = wb['日参数']
        h = self._header_map(ws)
        dc, tc = h.get('日期'), h.get('税比')
        if not dc or not tc:
            return
        for r in range(2, ws.max_row+1):
            d = _norm(ws.cell(row=r, column=dc).value)
            g = ws.cell(row=r, column=tc).value
            if d and g:
                self.tax_ratio[d] = float(g)

    def departments(self):
        return list(DEPTS)

    def has_date(self, date):
        return date in self.daily

    def _day_pretax(self, date):
        rec = self.daily.get(date)
        if rec is None:
            return None
        t = self.tax_ratio.get(date)
        out = {}
        for dept, v in rec.items():
            on, po = v['online_tax'], v['pos_tax']
            if t:
                out[dept] = {'online': on/t, 'pos': po/t, 'total': (on+po)/t}
            else:
                out[dept] = {'online': on, 'pos': po, 'total': on+po}
        return out

    def daily_pretax(self, date):
        return self._day_pretax(date) or {}

    def mtd_pretax(self, date):
        """本月1日至 date（含）累计未税预算。"""
        ym = date[:6]
        day = int(date[6:8])
        acc = {d: {'online':0.0,'pos':0.0,'total':0.0} for d in DEPTS}
        found_any = False
        for dd in range(1, day+1):
            dstr = f"{ym}{dd:02d}"
            one = self._day_pretax(dstr)
            if one is None:
                continue
            found_any = True
            for dept, v in one.items():
                for k in ('online','pos','total'):
                    acc[dept][k] += v[k]
        return acc if found_any else {}

    def daily_tax(self, date):
        """当日门店税比（含税/未税）。"""
        return self.tax_ratio.get(date)

    def mtd_tax(self, date):
        """本月1日至 date 的混合税比 = Σ含税预算 / Σ未税预算。"""
        ym = date[:6]
        day = int(date[6:8])
        s_tax = s_pretax = 0.0
        for dd in range(1, day+1):
            dstr = f"{ym}{dd:02d}"
            rec = self.daily.get(dstr)
            t = self.tax_ratio.get(dstr)
            if rec is None or not t:
                continue
            for v in rec.values():
                tot = v['online_tax'] + v['pos_tax']
                s_tax += tot
                s_pretax += tot / t
        return (s_tax / s_pretax) if s_pretax else None

    def min_date(self):
        return min(self.daily) if self.daily else None

    def max_date(self):
        return max(self.daily) if self.daily else None

class BudgetJsonReader:
    """从 budget_cat.json 读取预算，接口与 BudgetReader 完全兼容。

    JSON 结构: days[date] = {tax_ratio, depts:{dept:{online,pos,total}}, total}
    其中 online/pos/total 均为未税值（build_budget_lib.py 已做含税→未税转换）。
    """
    def __init__(self, path):
        import json
        with open(path, encoding='utf-8') as f:
            self.data = json.load(f)
        self.days = self.data.get('days', {})

    def departments(self):
        return list(DEPTS)

    def has_date(self, date):
        return date in self.days

    def _day(self, date):
        rec = self.days.get(date)
        if rec is None:
            return None
        return rec.get('depts', {})

    def daily_pretax(self, date):
        return dict(self._day(date)) if self._day(date) else {}

    def mtd_pretax(self, date):
        ym = date[:6]
        day = int(date[6:8])
        acc = {d: {'online': 0.0, 'pos': 0.0, 'total': 0.0} for d in DEPTS}
        found = False
        for dd in range(1, day + 1):
            dstr = f"{ym}{dd:02d}"
            one = self._day(dstr)
            if one is None:
                continue
            found = True
            for dept, v in one.items():
                if dept in acc:
                    for k in ('online', 'pos', 'total'):
                        acc[dept][k] += v.get(k, 0)
        return acc if found else {}

    def daily_tax(self, date):
        rec = self.days.get(date)
        return rec.get('tax_ratio') if rec else None

    def mtd_tax(self, date):
        ym = date[:6]
        day = int(date[6:8])
        s_tax = s_pretax = 0.0
        for dd in range(1, day + 1):
            dstr = f"{ym}{dd:02d}"
            rec = self.days.get(dstr)
            if rec is None:
                continue
            t = rec.get('tax_ratio')
            if not t:
                continue
            for v in rec.get('depts', {}).values():
                pretax = v.get('total', 0)
                s_tax += pretax * t
                s_pretax += pretax
        return (s_tax / s_pretax) if s_pretax else None

    def min_date(self):
        return min(self.days) if self.days else None

    def max_date(self):
        return max(self.days) if self.days else None


if __name__ == '__main__':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    p = sys.argv[1] if len(sys.argv) > 1 else r"D:\xiaoQ专用\Homes_MTD自动导出\预算分解.xlsx"
    br = BudgetReader(p)
    test = sys.argv[2] if len(sys.argv) > 2 else '20260907'
    print("有预算的日期数:", len(br.daily))
    d = br.daily_pretax(test)
    m = br.mtd_pretax(test)
    print(f"\n== {test} 单日未税预算 ==")
    for dept in DEPTS:
        print(f"  {dept}: 单日={d.get(dept,{}).get('total',0):.0f}  MTD累计={m.get(dept,{}).get('total',0):.0f}")
    print(f"  合计: 单日={sum(v.get('total',0) for v in d.values()):.0f}  MTD={sum(v.get('total',0) for v in m.values()):.0f}")
