#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
把《盒马Homes自动生成MTD营收对比表》xlsx 渲染成适合聊天窗口查看的长图 HTML。
57 列太宽，按 6 个逻辑区块（预算达成/整体成交/盒马APP/盒马POS/线上汇总/渠道占比）
纵向堆叠，每块保留“部门名称”列 + 总计行，宽度控制在 ~1100px 内，手机/聊天可直接读。
用法: python render_report_image.py <in.xlsx> <out.html>
"""
import sys, html
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as CIF

# ---- 列格式类型 ----
# pct=百分比(1位小数)  int=千分位整数  dec1=1位小数
def F(letter):
    return CIF(letter)

PCT = set("F G H I K L M O R T V W Z AA AC AE AH AI AK AM AN AO AQ AT AU AW AY AZ BA BB BC BD BE".split())
DEC1 = set("U AD AL AX".split())
# 其余数值列按整数千分位

def fmt(v, letter):
    """返回 (显示文本, 是否负数)"""
    if v is None:
        return ('', False)
    if isinstance(v, str):
        s = v.strip()
        if s in ('', '-'):
            return ('-', False)
        if s.startswith('#'):  # #DIV/0! 等
            return ('-', False)
        return (html.escape(s), False)
    # 数值
    neg = v < 0
    if letter in PCT:
        return (f"{v*100:.1f}%", neg)
    if letter in DEC1:
        return (f"{v:.1f}", neg)
    return (f"{round(v):,}", neg)

# ---- 区块定义: (区块名, [(列字母, 表头简称), ...]) ----
SECTIONS = [
    ("预算达成", [
        ("E","未税营业收入"), ("F","环比"), ("G","同比"), ("H","实际营收占比"),
        ("I","目标预算占比"), ("J","目标预算金额"), ("K","目标预算达成率"),
        ("L","损耗率"), ("M","管损率"),
    ]),
    ("整体成交", [
        ("N","成交金额"), ("O","环比"), ("P","日均成交金额"), ("Q","日均成交用户数"),
        ("R","环比"), ("S","日均门店成交用户数"), ("T","用户渗透率"),
        ("U","客单价"), ("V","客单价环比"), ("W","客单价同比"),
    ]),
    ("盒马APP", [
        ("X","成交金额"), ("Y","日均成交金额"), ("Z","环比"), ("AA","成交金额占比"),
        ("AB","日均成交用户数"), ("AC","环比"), ("AD","客单价"), ("AE","客单价环比"),
    ]),
    ("盒马POS", [
        ("AF","成交金额"), ("AG","日均成交金额"), ("AH","环比"), ("AI","成交金额占比"),
        ("AJ","日均成交用户数"), ("AK","环比"), ("AL","客单价"), ("AM","客单价环比"),
        ("AN","品类营收占比"), ("AO","目标预算占比"), ("AP","目标预算金额"), ("AQ","目标预算达成率"),
    ]),
    ("线上汇总", [
        ("AR","成交金额"), ("AS","日均成交金额"), ("AT","环比"), ("AU","成交金额占比"),
        ("AV","日均成交用户数"), ("AW","环比"), ("AX","客单价"), ("AY","客单价环比"),
        ("AZ","品类营收占比"),
    ]),
    ("渠道占比", [
        ("BA","盒马POS占比"), ("BB","盒马线上占比"), ("BC","盒马APP占比"),
        ("BD","盒马端外占比"), ("BE","环期占比"),
    ]),
]

# 区块标题配色
SEC_COLORS = ["1F4E5F", "2E6F4E", "8A4B08", "6B3FA0", "1F5AA8", "7A1F3D"]

def render(in_path, out_html):
    wb = load_workbook(in_path, data_only=True)
    ws = wb["Sheet1"]

    date = ws["A3"].value or ""
    store = ws["C3"].value or ""
    period_raw = str(ws["B3"].value or "").strip().lower()
    if period_raw == "1d":
        title_label, period_label = "单日", "单日 (1d)"
    else:
        title_label, period_label = "MTD", "月累计 (MTD)"
    # 数据行 3..17 为部门，18 为总计
    dept_rows = list(range(3, 18))
    total_row = 18

    def dept_name(r):
        d = ws.cell(row=r, column=4).value
        if d is None or str(d).strip() in ("", "-"):
            return "—"
        return str(d).strip()

    # 头部信息
    date_disp = str(date)
    if len(date_disp) == 8 and date_disp.isdigit():
        date_disp = f"{date_disp[:4]}-{date_disp[4:6]}-{date_disp[6:8]}"

    css = """
<style>
  * { margin:0; padding:0; box-sizing:border-box; }
  body { font-family:"Microsoft YaHei","Segoe UI",Arial,sans-serif; background:#F4F6F8; padding:20px; }
  .wrap { max-width:1120px; margin:0 auto; }
  .hero { background:linear-gradient(120deg,#1F4E5F,#2E6F4E); color:#fff; border-radius:12px;
          padding:20px 26px; margin-bottom:18px; }
  .hero h1 { font-size:22px; font-weight:700; letter-spacing:.5px; }
  .hero .sub { margin-top:8px; font-size:13px; opacity:.92; }
  .hero .sub b { font-weight:700; }
  .sec { background:#fff; border-radius:10px; overflow:hidden; margin-bottom:16px;
         box-shadow:0 1px 3px rgba(0,0,0,.08); }
  .sec-title { color:#fff; font-size:15px; font-weight:700; padding:9px 14px; letter-spacing:.5px; }
  table { width:100%; border-collapse:collapse; font-size:12px; }
  th,td { border:1px solid #D8DEE4; padding:4px 7px; white-space:nowrap; }
  thead th { background:#EAF1EE; color:#22313A; font-weight:700; text-align:center; }
  td.dept { text-align:left; font-weight:600; color:#22313A; background:#FBFCFD; }
  td.num { text-align:right; font-variant-numeric:tabular-nums; color:#22313A; }
  td.neg { color:#C0392B; font-weight:600; }
  tr.total td { background:#F2F6F4; font-weight:700; }
  tr.total td.dept { color:#1F4E5F; }
  thead th.dept-h { text-align:left; }
</style>
"""

    parts = ["<!DOCTYPE html><html lang='zh'><head><meta charset='utf-8'>", css, "</head><body><div class='wrap'>"]
    note = "　|　环比=比上周同日 · 单日无同比" if period_raw == "1d" else ""
    parts.append(f"<div class='hero'><h1>盒马 {title_label} 营收对比表 · {html.escape(store)}</h1>"
                 f"<div class='sub'>统计日期 <b>{date_disp}</b>　|　周期类型 <b>{period_label}</b>　|　共 {len(dept_rows)} 个部门 + 总计{note}</div></div>")

    for si, (sec_name, cols) in enumerate(SECTIONS):
        color = SEC_COLORS[si % len(SEC_COLORS)]
        parts.append(f"<div class='sec'><div class='sec-title' style='background:#{color};'>{html.escape(sec_name)}</div>")
        parts.append("<table><thead><tr><th class='dept-h'>部门名称</th>")
        for letter, head in cols:
            parts.append(f"<th>{html.escape(head)}</th>")
        parts.append("</tr></thead><tbody>")

        for r in dept_rows + [total_row]:
            is_total = (r == total_row)
            tr_cls = " class='total'" if is_total else ""
            parts.append(f"<tr{tr_cls}><td class='dept'>{html.escape(dept_name(r) if not is_total else '总计')}</td>")
            for letter, _ in cols:
                v = ws.cell(row=r, column=F(letter)).value
                txt, neg = fmt(v, letter)
                cls = "num neg" if neg else "num"
                parts.append(f"<td class='{cls}'>{txt}</td>")
            parts.append("</tr>")
        parts.append("</tbody></table></div>")

    parts.append("<div style='text-align:center;color:#8A97A0;font-size:11px;padding:6px 0 2px;'>"
                 "由 QwenWork 自动生成 · 数据源：盒马 Homes 店长大盘</div>")
    parts.append("</div></body></html>")

    with open(out_html, "w", encoding="utf-8") as f:
        f.write("".join(parts))
    print(f"HTML written: {out_html}")

def main():
    import io as _io
    sys.stdout = _io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    render(sys.argv[1], sys.argv[2])

if __name__ == "__main__":
    main()
